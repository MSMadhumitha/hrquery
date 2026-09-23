"""
Generates realistic, clearly-labeled-as-fictional sample HR policies for NovaTech Solutions.
Produces 8 documents across PDF, DOCX, and TXT formats with genuine sections, eligibility,
approval steps, and page structure.
"""

import sys
from pathlib import Path

# Add project root to sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import docx
from docx.shared import Inches, Pt, RGBColor
from config import SAMPLE_DOCS_DIR



def generate_pdf(filename: Path, title: str, pages_content: list[list[tuple[str, str]]]):
    """
    Builds a multi-page PDF document using ReportLab.
    pages_content is a list of pages, where each page is a list of (heading, paragraph_text) tuples.
    """
    doc = SimpleDocTemplate(
        str(filename),
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1A365D"),
        spaceAfter=12
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#2B6CB0"),
        spaceBefore=10,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#2D3748"),
        spaceAfter=8
    )
    notice_style = ParagraphStyle(
        'DocNotice',
        parent=styles['Italic'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#718096"),
        spaceAfter=15
    )
    
    story = []
    
    for page_idx, page_items in enumerate(pages_content):
        if page_idx == 0:
            story.append(Paragraph(title, title_style))
            story.append(Paragraph("<b>NovaTech Solutions Inc. — Internal Human Resources Policy Document [FICTIONAL SAMPLE]</b>", notice_style))
            story.append(Spacer(1, 10))
        else:
            story.append(Paragraph(f"{title} (Continued — Page {page_idx + 1})", h2_style))
            story.append(Spacer(1, 8))
            
        for heading, body in page_items:
            if heading:
                story.append(Paragraph(heading, h2_style))
            story.append(Paragraph(body, body_style))
            story.append(Spacer(1, 6))
            
        if page_idx < len(pages_content) - 1:
            story.append(PageBreak())
            
    doc.build(story)
    print(f"Generated PDF: {filename.name}")


def generate_docx(filename: Path, title: str, sections: list[tuple[str, str]]):
    """
    Builds a DOCX document using python-docx.
    """
    doc = docx.Document()
    
    # Title
    title_p = doc.add_heading(title, level=0)
    title_p.runs[0].font.color.rgb = RGBColor(26, 54, 93)
    
    # Fictional disclaimer
    sub = doc.add_paragraph("NovaTech Solutions Inc. — Internal HR Policy [FICTIONAL SAMPLE]")
    sub.runs[0].font.italic = True
    sub.runs[0].font.size = Pt(9)
    sub.runs[0].font.color.rgb = RGBColor(113, 128, 150)
    
    for heading, text in sections:
        if heading:
            h = doc.add_heading(heading, level=2)
            h.runs[0].font.color.rgb = RGBColor(43, 108, 176)
        p = doc.add_paragraph(text)
        p.style.font.size = Pt(10)
        p.style.font.color.rgb = RGBColor(45, 55, 72)
        
    doc.save(str(filename))
    print(f"Generated DOCX: {filename.name}")


def create_all_sample_documents():
    """Generates all 8 sample HR policy documents."""
    SAMPLE_DOCS_DIR.mkdir(parents=True, exist_ok=True)
    
    # 1. Employee Handbook (PDF - Multi-page)
    handbook_pages = [
        [
            ("1. Welcome and Company Philosophy",
             "Welcome to NovaTech Solutions Inc. Founded on innovation, integrity, and collaboration, NovaTech provides enterprise cloud and AI solutions globally. This Employee Handbook serves as a foundational guide to our standards, operating principles, and workplace practices."),
            ("2. Equal Employment Opportunity Policy",
             "NovaTech Solutions provides equal employment opportunities to all employees and applicants without regard to race, color, religion, gender, sexual orientation, gender identity, national origin, age, disability, or genetic information. Discrimination or harassment of any kind will not be tolerated."),
            ("3. Employment Classifications and Working Hours",
             "Standard operating hours are Monday through Friday, 9:00 AM to 6:00 PM local time. Regular full-time employees are scheduled for 40 hours per week. Flexible working hour arrangements must be agreed in writing with your direct manager.")
        ],
        [
            ("4. Probationary Period",
             "All newly hired employees undergo a mandatory 90-day probationary period. During this period, supervisors evaluate performance, aptitude, and cultural alignment. Successful completion is confirmed in writing by HR and management."),
            ("5. Annual Performance Reviews",
             "Performance reviews are conducted annually during the month of October. Employees complete a self-assessment, followed by supervisor evaluations and goal-setting sessions for the upcoming fiscal cycle. Performance ratings directly influence annual merit increments."),
            ("6. Workplace Safety and Health",
             "NovaTech maintains a zero-tolerance policy towards workplace hazards. Employees are required to report all accidents, injuries, and safety concerns immediately to the Workplace Safety Officer and Facilities Team.")
        ]
    ]
    generate_pdf(SAMPLE_DOCS_DIR / "Employee_Handbook.pdf", "NovaTech Employee Handbook", handbook_pages)
    
    # 2. Leave Policy (PDF - Multi-page)
    leave_pages = [
        [
            ("1. Purpose & Scope",
             "The NovaTech Leave Policy outlines the types of paid and unpaid leave available to full-time employees, accrual rules, and request mechanisms. All leaves must be logged through the NovaTech HR Portal."),
            ("2. Casual Leave (CL) Entitlement",
             "Full-time employees are entitled to 12 days of Casual Leave per calendar year, credited pro-rata on the first day of each quarter (3 days per quarter). Casual leave is intended for personal matters, emergencies, and brief personal absences. Employees must apply for casual leave at least 48 hours in advance through the HR Portal, subject to reporting manager approval. A maximum of 3 consecutive casual leaves may be taken at any one time. Casual leaves cannot be encashed or carried forward to the next calendar year; unused casual leaves lapse on December 31st.")
        ],
        [
            ("3. Sick Leave (SL)",
             "Employees receive 10 days of paid Sick Leave annually. For sick leave exceeding 2 consecutive working days, an official medical certificate signed by a licensed healthcare provider is mandatory upon returning to work."),
            ("4. Earned Leave / Privilege Leave (EL)",
             "Employees accrue 15 days of Earned Leave annually (1.25 days per completed month of active service). Earned leave requires 14 days prior notice and manager approval. Up to 30 days of unused earned leave may be carried forward into the subsequent calendar year."),
            ("5. Parental Leave",
             "Eligible female employees receive 26 weeks of fully paid Maternity Leave. Eligible male employees receive 2 weeks of fully paid Paternity Leave, to be availed within 6 months of childbirth or legal adoption.")
        ]
    ]
    generate_pdf(SAMPLE_DOCS_DIR / "Leave_Policy.pdf", "NovaTech Comprehensive Leave Policy", leave_pages)
    
    # 3. Attendance Policy (DOCX)
    attendance_sections = [
        ("1. Purpose and Overview",
         "NovaTech Solutions expects all employees to maintain reliable, punctual attendance to ensure seamless customer service and cross-functional team productivity."),
        ("2. Core Working Hours and Punctuality",
         "The standard core business hours are from 10:00 AM to 4:00 PM local time. All employees are required to be available and actively working during these core hours. NovaTech provides a 15-minute grace period for daily swipe-in (until 10:15 AM). More than three late arrivals beyond the grace period in a single month will be flagged for review by the reporting manager."),
        ("3. Half-Day Work Calculations",
         "An employee who logs between 4 and 6 hours of active work on a business day will be credited with a half-day's attendance. Any working duration below 4 hours without prior approved leave will be logged as an unexcused absence."),
        ("4. Unexcused Absences and Disciplinary Escalation",
         "Failure to report to work or notify the reporting supervisor for three (3) consecutive business days is classified as job abandonment and voluntary resignation, initiating formal HR review and potential termination of employment.")
    ]
    generate_docx(SAMPLE_DOCS_DIR / "Attendance_Policy.docx", "NovaTech Attendance and Punctuality Policy", attendance_sections)
    
    # 4. Work From Home Policy (PDF - Multi-page)
    wfh_pages = [
        [
            ("1. Policy Objective",
             "NovaTech embraces modern workplace flexibility while sustaining collaborative excellence. This Work From Home (WFH) Policy governs remote work arrangements, eligibility, and expectations."),
            ("2. Hybrid Work Model Eligibility",
             "Employees who have successfully completed their 90-day probationary period and whose job roles do not require physical on-premise presence are eligible for our hybrid model: up to two (2) days of remote work per calendar week, coordinated with team scheduling.")
        ],
        [
            ("3. Remote Work During Temporary Health Issues",
             "Can an employee work from home during a temporary health issue? Yes. When an employee experiences a temporary, non-critical medical condition, injury, or recovery phase that restricts commuting but does not impair the ability to perform core job functions, the employee may request temporary full-time remote work for up to four (4) consecutive weeks. To initiate this arrangement, the employee must submit a written request accompanied by a doctor's recommendation to both their reporting manager and the HR People Operations team. The request will be reviewed and approved within 48 hours."),
            ("4. Home Workspace and IT Security Requirements",
             "Employees working remotely must ensure a dedicated, ergonomic workspace with a secure high-speed internet connection. All work must be conducted on NovaTech-managed laptops connected via the corporate VPN. Company equipment stipends provide a one-time reimbursement of up to $300 for home office setup.")
        ]
    ]
    generate_pdf(SAMPLE_DOCS_DIR / "Work_From_Home_Policy.pdf", "NovaTech Work From Home (WFH) Policy", wfh_pages)
    
    # 5. Employee Benefits (DOCX)
    benefits_sections = [
        ("1. Comprehensive Employee Benefits Plan",
         "NovaTech Solutions offers a competitive benefits package designed to support physical, mental, and financial well-being for our employees and their families."),
        ("2. Group Health Insurance Coverage",
         "Full-time employees and their eligible dependents are covered under the NovaTech Premium Health Plan, providing up to $500,000 in comprehensive medical coverage, including inpatient hospitalization, prescription drugs, and preventive health screenings."),
        ("3. Wellness & Gym Reimbursement",
         "To foster active living, NovaTech provides a monthly wellness stipend of up to $50 per month ($600 annually). This stipend covers gym memberships, fitness class subscriptions, yoga studios, or authorized wellness apps. Receipts must be submitted quarterly."),
        ("4. 401(k) Retirement Savings Match",
         "NovaTech matches 100% of employee 401(k) contributions up to the first 5% of the employee's annual base salary. Company contributions vest immediately from day one of employment."),
        ("5. Continuous Learning & Development Stipend",
         "Each full-time team member is allocated an annual Professional Learning Allowance of $1,500 for technical certifications, conference attendances, and relevant university coursework approved by department leaders.")
    ]
    generate_docx(SAMPLE_DOCS_DIR / "Employee_Benefits.docx", "NovaTech Employee Benefits and Perks Guide", benefits_sections)
    
    # 6. Travel and Expense Policy (PDF - Multi-page)
    travel_pages = [
        [
            ("1. Overview and Pre-Travel Approvals",
             "This policy governs all business travel and legitimate out-of-pocket business expenses incurred by NovaTech employees. All travel requiring overnight lodging or airfare must receive pre-authorization from the departmental VP via the NovaTravel tool."),
            ("2. Flights and Lodging Guidelines",
             "Domestic flights under 6 continuous hours must be booked in Economy Class. International flights with continuous flying time exceeding 8 hours are eligible for Business Class booking. Hotel room rates are capped at $200 per night for standard tier cities and $300 per night for high-cost metropolitan areas (e.g., New York, London, San Francisco, Tokyo).")
        ],
        [
            ("3. Daily Meals and Per Diem Allowances",
             "NovaTech provides a fixed daily per diem for meals and incidental expenses: $75 per day for domestic travel and $120 per day for international destinations. Alcohol is excluded from company reimbursement."),
            ("4. Expense Submission and Reimbursement Process",
             "How are travel expenses reimbursed? Employees must submit all legitimate business travel expenses through the NovaExpense online portal within 15 calendar days of trip completion. Every expense item exceeding $25 must have an attached, clear, itemized digital receipt. Submissions are reviewed by the reporting manager and Finance department, and approved reimbursements are disbursed directly into the employee's bank account in the subsequent payroll cycle.")
        ]
    ]
    generate_pdf(SAMPLE_DOCS_DIR / "Travel_and_Expense_Policy.pdf", "NovaTech Business Travel and Expense Policy", travel_pages)
    
    # 7. Code of Conduct (PDF - Multi-page)
    conduct_pages = [
        [
            ("1. Ethical Foundations",
             "NovaTech Solutions holds all employees, contractors, and executives to the highest standards of professional honesty, fairness, and mutual respect in all business interactions."),
            ("2. Anti-Harassment and Non-Discrimination",
             "NovaTech strictly prohibits verbal, physical, sexual, or visual harassment. Any behavior creating an intimidating, hostile, or offensive working environment will result in immediate disciplinary intervention."),
            ("3. Conflict of Interest and Confidentiality",
             "Employees must not engage in outside business activities or secondary employment that competes with NovaTech or impairs their professional judgment. Proprietary source code, customer records, and trade secrets must remain confidential at all times.")
        ],
        [
            ("4. Disciplinary Consequences for Policy Violations",
             "What happens if an employee violates the company's code of conduct? Any reported violation is thoroughly investigated by the Ethics & Compliance Committee within ten (10) business days. Depending upon the gravity and intent of the infraction, NovaTech enforces progressive disciplinary actions: (a) First minor infraction: formal written warning and mandatory compliance counseling; (b) Second or repeated infraction: suspension without pay for up to fourteen (14) calendar days; (c) Serious violations—including theft, workplace violence, sexual harassment, fraud, or intentional breach of client confidentiality—result in immediate summary termination of employment for cause and potential legal referral."),
            ("5. Whistleblower Protection and Reporting Channels",
             "NovaTech guarantees strict confidentiality and zero retaliation for any employee who reports suspected ethical or legal misconduct in good faith via the anonymous Ethics Hotline.")
        ]
    ]
    generate_pdf(SAMPLE_DOCS_DIR / "Code_of_Conduct.pdf", "NovaTech Code of Business Conduct and Ethics", conduct_pages)
    
    # 8. Resignation Policy (TXT)
    resignation_content = """NOVATECH SOLUTIONS INC. — EMPLOYEE RESIGNATION AND EXIT POLICY
[FICTIONAL SAMPLE DOCUMENT FOR INTERNAL RAG TESTING]
Effective Date: January 1, 2024
Document ID: NTS-POL-HR-008-REV3

1. PURPOSE & SCOPE
This policy establishes the formal protocols and operational procedures governing voluntary employee resignations, notice periods, asset handovers, and full-and-final settlement at NovaTech Solutions Inc.

2. NOTICE PERIOD REQUIREMENTS
What is the resignation notice period?
- Individual Contributors (Engineers, Analysts, Designers, Coordinators): Required notice period is thirty (30) calendar days.
- People Managers, Technical Leads, and Project Managers: Required notice period is sixty (60) calendar days.
- Directors, Vice Presidents, and Executive Officers: Required notice period is ninety (90) calendar days.

The notice period commences on the business day on which formal written resignation is submitted via email to both the direct reporting manager and the Human Resources Department (hr-exits@novatech-fictional.com).

3. NOTICE PERIOD BUYOUT AND WAIVERS
Waiver of any portion of the contractual notice period is at the sole discretion of the Department Head and the VP of Human Resources. If an employee requests an early release and it is mutually agreed upon, a notice buyout calculation will apply based on base salary pro-rata.

4. ASSET RETURN AND KNOWLEDGE TRANSFER HANDOVER
Prior to the last working day (LWD), the departing employee must:
- Complete a comprehensive Knowledge Transfer (KT) document detailing current projects, credentials, and unresolved tasks.
- Return all NovaTech property, including company-issued laptops, access keycards, monitors, peripheral equipment, and corporate credit cards to the IT Asset Management department.
- Complete the Exit Clearance form signed off by IT, Finance, and Facilities.

5. FULL AND FINAL SETTLEMENT (F&F)
The full and final financial settlement, including encashment of eligible accumulated earned leaves, unpaid regular salary, and statutory deductions, will be finalized and transferred to the employee's registered bank account within thirty (30) calendar days following the official last working day.
"""
    with open(SAMPLE_DOCS_DIR / "Resignation_Policy.txt", "w", encoding="utf-8") as f:
        f.write(resignation_content)
    print("Generated TXT: Resignation_Policy.txt")
    print(f"All 8 sample documents successfully created in {SAMPLE_DOCS_DIR}")


if __name__ == "__main__":
    create_all_sample_documents()
