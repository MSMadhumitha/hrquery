"""
Document Loader module for HRQuery RAG.
Extracts structured text and metadata from PDF, DOCX, and TXT files.
Preserves page numbers (for PDFs) and source metadata.
"""

import io
import logging
from pathlib import Path
from typing import List, Dict, Any, Union, BinaryIO
import pypdf
import docx
from utils.helpers import clean_text

logger = logging.getLogger(__name__)


class DocumentLoader:
    """
    Loads and parses documents (PDF, DOCX, TXT), extracting clean text
    along with page numbers and document metadata.
    """

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

    @classmethod
    def load_document(
        cls, 
        source: Union[str, Path, BinaryIO], 
        filename: str = ""
    ) -> List[Dict[str, Any]]:
        """
        Parses a document file or byte stream.
        
        Returns a list of page/section dictionaries:
        [
            {
                "document_name": "Leave_Policy.pdf",
                "page_number": 1,
                "source_type": "pdf",
                "text": "..."
            }, ...
        ]
        """
        if isinstance(source, (str, Path)):
            path = Path(source)
            if not path.exists():
                raise FileNotFoundError(f"File not found: {path}")
            if not filename:
                filename = path.name
            extension = path.suffix.lower()
            with open(path, "rb") as f:
                content_bytes = f.read()
            stream = io.BytesIO(content_bytes)
        elif hasattr(source, "read"):
            # Handles BytesIO or Streamlit UploadedFile
            if not filename and hasattr(source, "name"):
                filename = getattr(source, "name")
            extension = Path(filename).suffix.lower()
            if hasattr(source, "seek"):
                source.seek(0)
            content_bytes = source.read()
            stream = io.BytesIO(content_bytes)
        else:
            raise ValueError("Unsupported source type. Must be a file path or file-like object.")

        if extension not in cls.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file format '{extension}' for file '{filename}'. "
                f"Supported formats: {', '.join(cls.SUPPORTED_EXTENSIONS)}"
            )

        if not content_bytes or len(content_bytes.strip()) == 0:
            logger.warning(f"File '{filename}' is empty.")
            return []

        if extension == ".pdf":
            return cls._load_pdf(stream, filename)
        elif extension == ".docx":
            return cls._load_docx(stream, filename)
        elif extension == ".txt":
            return cls._load_txt(content_bytes, filename)
        else:
            return []

    @classmethod
    def _load_pdf(cls, stream: io.BytesIO, filename: str) -> List[Dict[str, Any]]:
        """Extracts text page by page from PDF using pypdf."""
        results = []
        try:
            reader = pypdf.PdfReader(stream)
            num_pages = len(reader.pages)
            if num_pages == 0:
                logger.warning(f"PDF '{filename}' contains no pages.")
                return []

            for page_idx, page in enumerate(reader.pages, start=1):
                raw_text = page.extract_text() or ""
                text = clean_text(raw_text)
                if text:
                    results.append({
                        "document_name": filename,
                        "page_number": page_idx,
                        "source_type": "pdf",
                        "text": text
                    })
        except Exception as e:
            logger.error(f"Error parsing PDF '{filename}': {e}")
            raise ValueError(f"Corrupted or invalid PDF file '{filename}': {e}")

        return results

    @classmethod
    def _load_docx(cls, stream: io.BytesIO, filename: str) -> List[Dict[str, Any]]:
        """Extracts paragraphs and tables from DOCX using python-docx."""
        try:
            doc = docx.Document(stream)
            paragraphs = []
            
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    paragraphs.append(text)

            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            full_text = clean_text("\n\n".join(paragraphs))
            if not full_text:
                return []

            return [{
                "document_name": filename,
                "page_number": 1,
                "source_type": "docx",
                "text": full_text
            }]
        except Exception as e:
            logger.error(f"Error parsing DOCX '{filename}': {e}")
            raise ValueError(f"Corrupted or invalid DOCX file '{filename}': {e}")

    @classmethod
    def _load_txt(cls, content_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
        """Decodes text file with encoding fallback."""
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
        decoded_text = None
        for enc in encodings:
            try:
                decoded_text = content_bytes.decode(enc)
                break
            except (UnicodeDecodeError, LookupError):
                continue

        if decoded_text is None:
            raise ValueError(f"Unable to decode text file '{filename}' with supported encodings.")

        text = clean_text(decoded_text)
        if not text:
            return []

        return [{
            "document_name": filename,
            "page_number": 1,
            "source_type": "txt",
            "text": text
        }]
